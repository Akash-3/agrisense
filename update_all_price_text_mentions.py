import os
import shutil
import docx
from docx import Document

def update_synopsis_text_prices(docx_path):
    print(f"Updating all text price mentions in Synopsis {docx_path}...")
    doc = Document(docx_path)
    
    for p in doc.paragraphs:
        t = p.text
        if "5,930" in t or "6,000" in t:
            p.text = t.replace("5,930", "9,000").replace("6,000", "9,000")
            
    doc.save(docx_path)
    print(f"Done updating Synopsis {docx_path}!")

def update_report_text_prices(docx_path):
    print(f"Updating all text price mentions in Report {docx_path}...")
    doc = Document(docx_path)
    
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("2. Economic Feasibility:"):
            p.text = "2. Economic Feasibility: Commercial precision agriculture survey drones cost over ₹50,000 to ₹1,00,000, making them unaffordable for smallholder farmers. AgriSense constructs a complete dual-node IoT system for a total hardware budget of ₹9,000 by leveraging a Raspberry Pi Zero W & ESP32 boards (₹2,500), AS7341 optical sensor (₹2,500), soil/climate/gas probes (₹1,500), and drone frame & power (₹2,500), combined with 100% free open-source software (FastAPI, React, PyTorch). (See Section 4.b for itemized breakdown)."
        elif "5,930" in t or "6,000" in t:
            p.text = t.replace("5,930", "9,000").replace("6,000", "9,000")

    doc.save(docx_path)
    print(f"Done updating Report {docx_path}!")

if __name__ == "__main__":
    docs_dir = r"C:\Users\tempm\.gemini\antigravity\scratch\agrisense\docs"
    
    syn1 = os.path.join(docs_dir, "AgriSense_Project_Synopsis.docx")
    syn2 = r"C:\Users\tempm\OneDrive\Documents\AgriSense_Project_Synopsis.docx"
    update_synopsis_text_prices(syn1)
    update_synopsis_text_prices(syn2)

    rep1 = r"C:\Users\tempm\OneDrive\Documents\AgriSense Report.docx"
    rep2 = os.path.join(docs_dir, "AgriSense_Final_Report.docx")
    update_report_text_prices(rep1)
    update_report_text_prices(rep2)

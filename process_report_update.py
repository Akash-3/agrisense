import os
import shutil
import glob
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def add_page_number_to_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    # Add text "Page "
    run1 = p.add_run("Page ")
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(10)
    run1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Add PAGE field
    fldSimple1 = OxmlElement('w:fldSimple')
    fldSimple1.set(qn('w:instr'), 'PAGE')
    p._element.append(fldSimple1)
    
    # Add text " of "
    run2 = p.add_run(" of ")
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Add NUMPAGES field
    fldSimple2 = OxmlElement('w:fldSimple')
    fldSimple2.set(qn('w:instr'), 'NUMPAGES')
    p._element.append(fldSimple2)

def make_backups(src_path):
    print("Making backups of the original document...")
    backup_onedrive = src_path.replace(".docx", "_BACKUP.docx")
    backup_scratch = r"C:\Users\tempm\.gemini\antigravity\scratch\agrisense\docs\AgriSense_Report_BACKUP.docx"
    
    shutil.copy(src_path, backup_onedrive)
    print(f"Backup created at: {backup_onedrive}")
    
    shutil.copy(src_path, backup_scratch)
    print(f"Backup created at: {backup_scratch}")
    return backup_onedrive

def create_gantt_table(doc):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("Project Implementation Schedule & Gantt Chart (Red Timeline)")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    table = doc.add_table(rows=15, cols=13)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Activity / Task", "W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8", "W9", "W10", "W11", "W12"]
    for i, title in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = title
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(c, "1E3A8A")

    # Activities and active week ranges (1-indexed weeks)
    activities = [
        ("Literature Review", [1]),
        ("Requirements Analysis", [2]),
        ("Hardware Procurement", [3, 4]),
        ("ESP32 Development", [3, 4]),
        ("Sensor Calibration", [5, 6]),
        ("Data Collection", [6, 7]),
        ("Backend Development", [7, 8]),
        ("Database Integration", [8, 9]),
        ("AI Model Development", [9, 10]),
        ("AI Training", [9, 10]),
        ("React Dashboard", [10, 11]),
        ("UI & Charts", [11]),
        ("System Testing", [11, 12]),
        ("Documentation", [12])
    ]

    for row_idx, (act_name, active_weeks) in enumerate(activities, start=1):
        row_cells = table.rows[row_idx].cells
        row_cells[0].text = act_name
        p_act = row_cells[0].paragraphs[0]
        if p_act.runs:
            p_act.runs[0].font.size = Pt(9)
            p_act.runs[0].font.bold = True
        
        for w in range(1, 13):
            c_week = row_cells[w]
            p_w = c_week.paragraphs[0]
            p_w.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            if w in active_weeks:
                # Red line timeline fill (#DC2626)
                set_cell_background(c_week, "DC2626")
                r = p_w.add_run("■")
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                set_cell_background(c_week, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(12)
    r_n = p_note.add_run("Table 2.1: AgriSense 12-Week Implementation Gantt Chart (Red Bar Execution Timeline).")
    r_n.font.size = Pt(9.5)
    r_n.font.italic = True
    r_n.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

if __name__ == "__main__":
    src_docx = r"C:\Users\tempm\OneDrive\Documents\AgriSense Report.docx"
    make_backups(src_docx)

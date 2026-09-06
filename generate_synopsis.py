import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def generate_synopsis_doc(docx_path):
    doc = Document()
    
    # Margins & Base Page Setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        return p

    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        return p

    def add_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.bold = True
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.bold = True
        p.add_run(text)
        return p

    # Title Header
    add_title("PROJECT SYNOPSIS: AGRISENSE")
    add_subtitle("A Smart Flying Drone & Soil Sensor Helper for Farmers with AI Disease Warning")
    
    # 1. Introduction
    add_heading("1. Project Introduction (What is AgriSense?)")
    add_p("Imagine having a smart flying helper (a drone) and a ground helper (a soil sensor) that work together to take care of farm plants. That is what AgriSense is!")
    add_p("Plants can get sick from tiny fungi or lack of water. Normally, by the time a farmer sees brown spots on leaves, the plant is already very sick. AgriSense uses a special 10-color light sensor on a drone to look at leaves in a way human eyes cannot. It can see inside the leaf to spot sickness 5 days early! At the same time, a sensor in the dirt checks if the soil is thirsty. All this information goes to a simple website that shows the farmer if their crops are happy and healthy.")

    # 2. Problem Statement
    add_heading("2. The Problem (Why Farmers Need Help)")
    add_p("Farmers face major challenges that make growing food difficult:")
    add_bullet("By the time brown spots or yellow leaves show up, the disease has already spread inside the plant. Medicines applied this late do not work well.", bold_prefix="• Sick Leaves Show Symptoms Too Late:")
    add_bullet("Walking around huge fields to check every single plant takes days and is very tiring, so many sick plants get missed.", bold_prefix="• Fields Are Too Big to Check by Hand:")
    add_bullet("Without knowing exact soil moisture, farmers might give plants too much water or too little water.", bold_prefix="• Hard to Guess Water Needs:")
    add_bullet("Expensive laboratory cameras cost thousands of dollars, which small farmers cannot afford.", bold_prefix="• Professional Tools Cost Too Much:")

    # 3. Objectives & Features
    add_heading("3. Project Objectives & Key Features")
    add_p("AgriSense is built with five simple, powerful goals in mind:")
    add_bullet("Detect plant sickness 5 days before human eyes can see brown spots.", bold_prefix="1. Early Sickness Warning:")
    add_bullet("Measure soil water levels continuously so crops get the exact right amount of water.", bold_prefix="2. Smart Water Tracking:")
    add_bullet("Keep the drone light (only 82 grams!) so it flies easily and stays in the air longer.", bold_prefix="3. Feather-Light Drone Flying:")
    add_bullet("Send warning alerts if the field gets too hot, dry, or if there is smoke/fire nearby.", bold_prefix="4. Emergency Fire & Hazard Alert:")
    add_bullet("Show everything on a simple, colorful web dashboard with clear green, yellow, and red badges.", bold_prefix="5. Kid-Simple Web Dashboard:")

    # 4. Methodology
    add_heading("4. How It Works (Methodology Step-by-Step)")
    add_p("The project works in four easy steps:")
    add_bullet("The drone flies over crop fields. Its 10-color light sensor (AS7341) measures how light reflects off leaves, including invisible Near-Infrared light. A sensor in the dirt checks soil moisture.", bold_prefix="Step 1 - Collect Data:")
    add_bullet("The small computer chip (ESP32) packs all numbers into a tiny digital message and sends it through Wi-Fi to our backend computer.", bold_prefix="Step 2 - Send Wirelessly:")
    add_bullet("Our artificial intelligence (AI) brain, called MM-SSNet, reads the light colors. It compares leaf greenness and internal leaf health to spot sickness early.", bold_prefix="Step 3 - Think with AI:")
    add_bullet("The website updates every 5 seconds. Green means Healthy, Yellow means Caution, and Red means Warning!", bold_prefix="Step 4 - Show Results:")

    # 5. Diagram & Flowchart
    add_heading("5. System Architecture & Flowchart (Draw.io Diagram)")
    add_p("Here is the visual map of how AgriSense works, modeled in Draw.io (AgriSense_Architecture.drawio):")
    
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["System Component", "Simple Role / What it Does"]
    for i, title in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = title
        c.paragraphs[0].runs[0].font.bold = True
        set_cell_background(c, "1E3A8A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    components = [
        ("Flying Drone Payload Node", "Carries ESP32 chip [1] + 10-color sensor [2] + Temp/Smoke sensors. Flies over crops to scan leaves."),
        ("Ground Soil Node", "Stays in the dirt with a capacitive moisture probe [12]. Measures how wet or dry the soil is."),
        ("FastAPI AI Computer Server", "Receives data via Wi-Fi [3]. Uses MM-SSNet AI [15] to check plant health under 200 milliseconds."),
        ("React Web Dashboard", "Shows real-time colorful charts [4, 5] and alerts on phones or computers so farmers know what to do.")
    ]
    for r_idx, (comp, desc) in enumerate(components):
        row_cells = table.rows[r_idx+1].cells
        row_cells[0].text = comp
        row_cells[1].text = desc

    add_p("\nFlow Sequence: Drone Scans Leaves -> Wi-Fi Sends Data -> AI Server Checks Health -> Web Dashboard Shows Results!")

    # 6. Software & Hardware Requirements
    add_heading("6. Hardware and Software Requirements")
    add_p("To build AgriSense, we use friendly, accessible tools:")
    add_bullet("ESP32 DevKit V1 Microcontroller (Tiny $4 computer chip with Wi-Fi) [1]", bold_prefix="• Hardware 1:")
    add_bullet("Adafruit AS7341 10-Channel Multi-Spectral Light Sensor (Reads 10 colors including NIR) [2]", bold_prefix="• Hardware 2:")
    add_bullet("Capacitive Soil Moisture Sensor v1.2 (Rust-free soil probe) [12]", bold_prefix="• Hardware 3:")
    add_bullet("DHT22 Air Temperature & Humidity Sensor + MQ-2 Smoke/Gas Sensor", bold_prefix="• Hardware 4:")
    add_bullet("Python FastAPI (Super fast web server engine) [3]", bold_prefix="• Software 1:")
    add_bullet("MM-SSNet Artificial Intelligence Neural Network (PyTorch / MobileNetV3) [15]", bold_prefix="• Software 2:")
    add_bullet("React.js & Chart.js Web Dashboard (Modern web frontend) [4, 5]", bold_prefix="• Software 3:")
    add_bullet("PostgreSQL Database (Remembers all past measurements)", bold_prefix="• Software 4:")

    # 7. Cost Estimation
    add_heading("7. Budget & Cost Estimation")
    add_p("AgriSense is engineered to be super affordable so any small farmer or school lab can make one. Industrial farm drones cost over $5,000 to $10,000, but AgriSense is built entirely from pocket-friendly hobby electronics:")
    add_bullet("ESP32 Microcontroller Boards (2 units for drone & ground) ~ Low Cost", bold_prefix="1. Core Chips:")
    add_bullet("Adafruit AS7341 Spectral Sensor + Soil Probe + Climate Sensors ~ Low Cost", bold_prefix="2. Sensors:")
    add_bullet("Makeshift Drone Frame + Small LiPo Battery ~ Low Cost", bold_prefix="3. Drone Frame:")
    add_bullet("FastAPI, React.js, Python, PostgreSQL, Chart.js ~ 100% FREE Open Source!", bold_prefix="4. Software & AI:")
    add_p("Total Estimated Hardware Budget: Extremely low-cost and budget-friendly, making it 100x cheaper than commercial agricultural systems!")

    # 8. Conclusion
    add_heading("8. Conclusion")
    add_p("AgriSense proves that we do not need expensive industrial equipment to help farmers grow healthy crops. By combining a smart flying drone, a ground soil sensor, and an AI brain, AgriSense catches plant diseases 5 days early and saves water. It is fast, affordable, easy to use, and helps keep our food supply safe and strong!")

    doc.save(docx_path)
    print(f"Successfully generated Synopsis Docx at: {docx_path}")

if __name__ == "__main__":
    docs_dir = r"C:\Users\tempm\.gemini\antigravity\scratch\agrisense\docs"
    synopsis_docx = os.path.join(docs_dir, "AgriSense_Project_Synopsis.docx")
    synopsis_md = os.path.join(docs_dir, "AgriSense_Project_Synopsis.md")
    
    generate_synopsis_doc(synopsis_docx)
    
    # Export Markdown
    doc = Document(synopsis_docx)
    md_lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if t.startswith("PROJECT SYNOPSIS"):
            md_lines.append(f"# {t}\n")
        elif t.startswith("1.") or t.startswith("2.") or t.startswith("3.") or t.startswith("4.") or t.startswith("5.") or t.startswith("6.") or t.startswith("7.") or t.startswith("8."):
            md_lines.append(f"\n## {t}\n")
        else:
            md_lines.append(f"{t}\n")
            
    with open(synopsis_md, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    print(f"Successfully generated Synopsis MD at: {synopsis_md}")

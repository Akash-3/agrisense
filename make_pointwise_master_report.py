import os
import docx
from docx import Document
from docx.oxml import OxmlElement

def add_paragraph_after(paragraph, text, style='List Bullet'):
    new_p = OxmlElement('w:p')
    paragraph._element.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    new_para.style = style
    new_para.add_run(text)
    return new_para

def transform_master_report_to_pointwise(docx_path):
    print(f"Modifying {docx_path} IN-PLACE to add point-wise formatting...")
    doc = Document(docx_path)
    
    for p in doc.paragraphs:
        t = p.text.strip()
        
        # Abstract
        if t.startswith("Agricultural yield loss attributed to unpredictable phytopathological outbreaks"):
            p.text = "Abstract Key Background Points:"
            p.style = 'Normal'
            add_paragraph_after(p, "Fungal and bacterial pathogen outbreaks account for 20% to 40% of annual global crop destruction.", style='List Bullet')
            add_paragraph_after(p, "Standard 2D RGB camera vision systems (ResNet, YOLO) are strictly reactive, spotting disease only after macroscopic leaf lesions physically emerge.", style='List Bullet')
            add_paragraph_after(p, "Laboratory-grade hyperspectral cameras cost thousands of dollars, making them prohibitively expensive for smallholder farmers.", style='List Bullet')

        elif t.startswith("This project presents AgriSense, an end-to-end, low-cost Internet of Things"):
            p.text = "AgriSense Solution & Sensing Architecture:"
            p.style = 'Normal'
            add_paragraph_after(p, "Decoupled Aerial & Ground Sensing: Aerial canopy scanning is separated from terrestrial ground soil probing to prevent drone weight strain.", style='List Bullet')
            add_paragraph_after(p, "Lightweight Drone Sensor Payload: Weighs 82 grams, incorporating an ESP32 DevKit V1 [1], Adafruit AS7341 10-channel spectral sensor [2], DHT22 microclimate probe, and MQ-2 smoke sensor.", style='List Bullet')
            add_paragraph_after(p, "Stationary Terrestrial Ground Nodes: Capacitive soil moisture probes [12] sit in the dirt to measure dielectric permittivity without metal corrosion.", style='List Bullet')

        # Chapter 1.a Objective & Scope
        elif t.startswith("Precision agriculture represents the convergence of information technology"):
            p.text = "Precision Agriculture & Global Agricultural Challenges:"
            p.style = 'Normal'
            add_paragraph_after(p, "Global Yield Protection: According to the FAO, plant diseases cause hundreds of billions of dollars in annual economic losses.", style='List Bullet')
            add_paragraph_after(p, "Smallholder Resource Bottlenecks: Small farmers lack affordable access to plant agronomists or expensive satellite systems.", style='List Bullet')
            add_paragraph_after(p, "Reactive Intervention Limits: Current disease intervention occurs only after extensive foliage browning has already reduced yield.", style='List Bullet')

        # Chapter 1.b Theoretical Background
        elif t.startswith("Understanding plant canopy optics requires analyzing how vegetative tissues"):
            p.text = "Plant Canopy Optical Physics:"
            p.style = 'Normal'
            add_paragraph_after(p, "Photosynthetic Pigment Absorption: Chlorophyll-a and chlorophyll-b strongly absorb blue (400-500nm) and red (620-680nm) light while reflecting green light (520-560nm).", style='List Bullet')
            add_paragraph_after(p, "Near-Infrared Mesophyll Scattering: Healthy hydrated leaf spongy mesophyll cell walls scatter up to 50% of incident Near-Infrared light (700-900nm).", style='List Bullet')
            add_paragraph_after(p, "Pre-Symptomatic Structural Degradation: Pathogen hyphae break down cell turgor pressure, causing NIR scattering to collapse 5.4 days before visible leaf chlorosis emerges.", style='List Bullet')

    doc.save(docx_path)
    print(f"Successfully updated {docx_path} IN-PLACE with point-wise bullet lists!")

if __name__ == "__main__":
    master_path = r"C:\Users\tempm\.gemini\antigravity\scratch\agrisense\docs\AgriSense_Final_Report.docx"
    master_md = r"C:\Users\tempm\.gemini\antigravity\scratch\agrisense\docs\AgriSense_Final_Report.md"
    
    transform_master_report_to_pointwise(master_path)
    
    # Export Markdown
    doc = Document(master_path)
    md_lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if t.startswith("CHAPTER") or t.startswith("AGRISENSE:") or t.startswith("CERTIFICATE") or t.startswith("DECLARATION") or t.startswith("ACKNOWLEDGEMENTS") or t.startswith("ABSTRACT") or t.startswith("TABLE OF CONTENTS") or t.startswith("BIBLIOGRAPHY"):
            md_lines.append(f"\n# {t}\n")
        elif t.startswith("1.") or t.startswith("2.") or t.startswith("3.") or t.startswith("4.") or t.startswith("5.") or t.startswith("6.") or t.startswith("7.") or t.startswith("8.") or t.startswith("9."):
            md_lines.append(f"\n## {t}\n")
        elif p.style.name == 'List Bullet':
            md_lines.append(f"* {t}\n")
        elif "Figure " in t:
            if "Figure 3.1" in t:
                md_lines.append(f"\n![Aerial Drone Payload](agrisense_drone_payload.jpg)\n*{t}*\n")
            elif "Figure 3.2" in t:
                md_lines.append(f"\n![Ground Node Unit](agrisense_ground_node.jpg)\n*{t}*\n")
            else:
                md_lines.append(f"*{t}*\n")
        else:
            md_lines.append(f"{t}\n")
            
    with open(master_md, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    print(f"Successfully updated {master_md}!")

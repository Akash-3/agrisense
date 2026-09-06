import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_cover_page_to_existing_file(docx_path):
    print(f"Adding formal cover page to {docx_path} without changing anything else...")
    
    # First make a safety backup of Downloads file
    backup_path = docx_path.replace(".docx", "_BACKUP.docx")
    shutil.copy(docx_path, backup_path)
    print(f"Created safety backup at: {backup_path}")

    doc = Document(docx_path)
    
    # Find the very first paragraph in the document
    first_p = doc.paragraphs[0]
    
    # Build cover page items to insert before first_p
    # Note: To insert in correct order before first_p, we can create list of paragraphs and insert them sequentially
    
    cover_elements = [
        ("title", "AgriSense: A Low-Cost IoT Drone Platform for Early Crop Disease Diagnosis and Soil Health Monitoring Using Dual-Stream Spectral-Spatial AI"),
        ("subtitle", "A Project Synopsis Submitted in Partial Fulfillment of the Requirements for T.Y. B.Sc. IT Semester-V"),
        ("submitted_by", "Submitted By:"),
        ("group", "Group – G580"),
        ("author1", "Kavya Bhandary, 261796"),
        ("author2", "Rehaan Shaikh, 261831"),
        ("author3", "Aditya Surve, 261844"),
        ("guidance_by", "Under the Guidance of:"),
        ("guide", "Jinal Gujar"),
        ("dept", "DEPARTMENT OF INFORMATION TECHNOLOGY"),
        ("year", "ACADEMIC YEAR 2025 – 2026")
    ]

    for elem_type, text in cover_elements:
        p = first_p.insert_paragraph_before()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if elem_type == "title":
            p.paragraph_format.space_before = Pt(36)
            p.paragraph_format.space_after = Pt(18)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(22)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            
        elif elem_type == "subtitle":
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(64)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        elif elem_type == "submitted_by":
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        elif elem_type == "group":
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        elif elem_type.startswith("author"):
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        elif elem_type == "guidance_by":
            p.paragraph_format.space_before = Pt(28)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        elif elem_type == "guide":
            p.paragraph_format.space_after = Pt(48)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        elif elem_type == "dept":
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        elif elem_type == "year":
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Insert page break right before first_p
    p_break = first_p.insert_paragraph_before()
    p_break.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

    doc.save(docx_path)
    print(f"Successfully added formal cover page to {docx_path}!")

if __name__ == "__main__":
    target_docx = r"C:\Users\tempm\Downloads\AgriSense_Project_Synopsis.docx"
    add_cover_page_to_existing_file(target_docx)
